from app import marca
import re
import io
import os
from datetime import date, datetime, timedelta, timezone
from flask import Blueprint, render_template, redirect, url_for, request, flash, send_file, current_app, jsonify
from flask_login import login_required, current_user
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ... import db
from ...models import (Pessoa, Atendimento, Solicitacao, LogAcao, TagInteresse,
                       RequerimentoGerado, ATENDIMENTO_CONCLUIDO)

pessoas_bp = Blueprint('pessoas', __name__, url_prefix='/pessoas')


def _registrar_log(acao, entidade, entidade_id, descricao):
    """Registra uma ação crítica no log."""
    log = LogAcao(
        id_usuario=current_user.id if current_user.is_authenticated else None,
        acao=acao,
        entidade=entidade,
        entidade_id=entidade_id,
        descricao=descricao,
    )
    db.session.add(log)

POR_PAGINA = 20

BAIRROS_MUNICIPIO = list(marca.BAIRROS)

_CPF_RE  = re.compile(r'^\d{3}\.\d{3}\.\d{3}-\d{2}$')
_CNPJ_RE = re.compile(r'^\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}$')


def _cpf_valido(cpf: str) -> bool:
    """Valida CPF pelos dígitos verificadores."""
    digitos = ''.join(filter(str.isdigit, cpf))
    if len(digitos) != 11 or len(set(digitos)) == 1:
        return False
    soma = sum(int(digitos[i]) * (10 - i) for i in range(9))
    d1 = (soma * 10 % 11) % 10
    if d1 != int(digitos[9]):
        return False
    soma = sum(int(digitos[i]) * (11 - i) for i in range(10))
    d2 = (soma * 10 % 11) % 10
    return d2 == int(digitos[10])


def _cnpj_valido(cnpj: str) -> bool:
    """Valida CNPJ pelos dígitos verificadores."""
    digitos = ''.join(filter(str.isdigit, cnpj))
    if len(digitos) != 14 or len(set(digitos)) == 1:
        return False
    pesos1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    pesos2 = [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    soma = sum(int(digitos[i]) * pesos1[i] for i in range(12))
    d1 = 0 if soma % 11 < 2 else 11 - soma % 11
    if d1 != int(digitos[12]):
        return False
    soma = sum(int(digitos[i]) * pesos2[i] for i in range(13))
    d2 = 0 if soma % 11 < 2 else 11 - soma % 11
    return d2 == int(digitos[13])


def _valida_form(form):
    """Valida os campos do formulário. Retorna lista de erros."""
    erros = []

    nome = form.get('nome_razao_social', '').strip()
    if not nome:
        erros.append('Nome / Razão Social é obrigatório.')

    cpf = form.get('cpf', '').strip()
    if cpf and not _CPF_RE.match(cpf):
        erros.append('CPF deve estar no formato 000.000.000-00.')
    if cpf and _CPF_RE.match(cpf) and not _cpf_valido(cpf):
        erros.append('CPF inválido — verifique os dígitos.')

    cnpj = form.get('cnpj', '').strip()
    if cnpj and not _CNPJ_RE.match(cnpj):
        erros.append('CNPJ deve estar no formato 00.000.000/0000-00.')
    if cnpj and _CNPJ_RE.match(cnpj) and not _cnpj_valido(cnpj):
        erros.append('CNPJ inválido — verifique os dígitos.')

    # Se não tem CNPJ (não é PJ), CPF é obrigatório
    if not cnpj and not cpf:
        erros.append('CPF é obrigatório para Pessoa Física.')

    data_nasc = form.get('data_nascimento', '').strip()
    if data_nasc:
        try:
            date.fromisoformat(data_nasc)
        except ValueError:
            erros.append('Data de nascimento inválida.')

    return erros


def _salva_pessoa(pessoa, form):
    """Aplica os dados do form em um objeto Pessoa."""
    pessoa.nome_razao_social  = form.get('nome_razao_social', '').strip()
    pessoa.email              = form.get('email',     '').strip() or None
    pessoa.cpf                = form.get('cpf',       '').strip() or None
    pessoa.cnpj               = form.get('cnpj',      '').strip() or None
    pessoa.telefone           = form.get('telefone',  '').strip() or None
    pessoa.endereco     = form.get('endereco',     '').strip() or None
    pessoa.cep          = form.get('cep',          '').strip() or None
    pessoa.logradouro   = form.get('logradouro',   '').strip() or None
    pessoa.numero       = form.get('numero',       '').strip() or None
    pessoa.complemento  = form.get('complemento',  '').strip() or None
    pessoa.bairro       = form.get('bairro',       '').strip() or None
    pessoa.sexo         = form.get('sexo',         '').strip() or None

    # Dados Receita Federal
    pessoa.nome_fantasia      = form.get('nome_fantasia',      '').strip() or None
    pessoa.situacao_cadastral = form.get('situacao_cadastral', '').strip() or None
    pessoa.capital_social     = form.get('capital_social',     '').strip() or None
    pessoa.porte              = form.get('porte',              '').strip() or None
    pessoa.natureza_juridica  = form.get('natureza_juridica',  '').strip() or None
    pessoa.cnae_principal     = form.get('cnae_principal',     '').strip() or None
    pessoa.cnae_secundario    = form.get('cnae_secundario',    '').strip() or None
    pessoa.inscricao_estadual = form.get('inscricao_estadual', '').strip() or None

    # Ficha física — dados complementares
    pessoa.rg                  = form.get('rg',                  '').strip() or None
    pessoa.nivel_escolar       = form.get('nivel_escolar',       '').strip() or None
    pessoa.renda_familiar      = form.get('renda_familiar',      '').strip() or None
    pessoa.instagram           = form.get('instagram',           '').strip() or None
    pessoa.endereco_comercial  = form.get('endereco_comercial',  '').strip() or None
    pessoa.principais_produtos = form.get('principais_produtos', '').strip() or None
    pessoa.faturamento_mensal  = form.get('faturamento_mensal',  '').strip() or None
    pessoa.observacoes_cadastro = form.get('observacoes_cadastro', '').strip() or None

    nmf = form.get('num_membros_familia', '').strip()
    pessoa.num_membros_familia = int(nmf) if nmf.isdigit() else None
    nf = form.get('num_funcionarios', '').strip()
    pessoa.num_funcionarios = int(nf) if nf.isdigit() else None

    mei = form.get('mei')
    pessoa.mei = (mei in ('true', '1', True)) if mei is not None else None

    simples = form.get('simples_nacional')
    pessoa.simples_nacional = (simples in ('true', '1', True)) if simples is not None else None

    for campo, key in [
        ('data_nascimento',        'data_nascimento'),
        ('data_situacao_cadastral','data_situacao_cadastral'),
        ('data_inicio_atividade',  'data_inicio_atividade'),
    ]:
        val = form.get(key, '').strip()
        setattr(pessoa, campo, date.fromisoformat(val) if val else None)


# ── Lista ──────────────────────────────────────────────────────────────────────
@pessoas_bp.route('/')
@login_required
def index():
    q = request.args.get('q', '').strip()
    query = Pessoa.visiveis()

    if q:
        like = f'%{q}%'
        q_digitos = re.sub(r'\D', '', q)
        filtros = [
            Pessoa.nome_razao_social.ilike(like),
            Pessoa.cpf.ilike(like),
            Pessoa.cnpj.ilike(like),
            Pessoa.telefone.ilike(like),
            Pessoa.nome_fantasia.ilike(like),
        ]
        if q_digitos:
            like_d = f'%{q_digitos}%'
            filtros += [
                db.func.regexp_replace(Pessoa.cpf,  r'\D', '', 'g').ilike(like_d),
                db.func.regexp_replace(Pessoa.cnpj, r'\D', '', 'g').ilike(like_d),
            ]
        query = query.filter(db.or_(*filtros))

    if current_user.perfil.nome != 'administrador':
        query = query.filter(Pessoa.exclusao_pendente == False)

    # ── Filtros rápidos (chips) ────────────────────────────────────────────
    filtro_tipo     = request.args.get('tipo', '')
    filtro_especial = request.args.get('especial', '')
    filtro_tag      = request.args.get('tag', type=int)

    if filtro_tipo == 'PF':
        # Pessoa física pura — tem CPF mas não tem CNPJ
        query = query.filter(
            Pessoa.cpf != None,
            Pessoa.cpf != '',
            db.or_(Pessoa.cnpj == None, Pessoa.cnpj == '')
        )
    elif filtro_tipo == 'MEI':
        # MEI — flag mei explicitamente True (None/False excluídos)
        query = query.filter(Pessoa.mei == True)
    elif filtro_tipo == 'PJ':
        # Pessoa jurídica não MEI — tem CNPJ mas mei não é True
        query = query.filter(
            Pessoa.cnpj != None,
            Pessoa.cnpj != '',
            db.or_(Pessoa.mei == None, Pessoa.mei == False)
        )

    if filtro_especial == 'sem_telefone':
        query = query.filter(
            db.or_(Pessoa.telefone == None, Pessoa.telefone == '')
        )
    elif filtro_especial == 'sem_atendimento':
        recentes = db.session.query(Atendimento.id_pessoa).filter(
            Atendimento.data_atendimento >= date.today() - timedelta(days=90),
            Atendimento.status == ATENDIMENTO_CONCLUIDO
        ).subquery()
        query = query.filter(~Pessoa.id.in_(recentes))

    if filtro_tag:
        query = query.filter(Pessoa.tags.any(TagInteresse.id == filtro_tag))

    page = request.args.get('page', 1, type=int)
    paginacao = (
        query
        .order_by(Pessoa.nome_razao_social)
        .paginate(page=page, per_page=POR_PAGINA, error_out=False)
    )

    # ── Badge documentação incompleta ─────────────────────────────────────
    _tipos_obrig = {'rg_cnh', 'comprovante_endereco'}
    ids_doc_incompleta = {
        p.id for p in paginacao.items
        if not _tipos_obrig.issubset({d.tipo for d in p.documentos_visiveis if d.ativo})
    }

    tags_disponiveis = TagInteresse.query.filter_by(ativo=True).order_by(TagInteresse.descricao).all()

    from datetime import date as _date
    return render_template('pessoas/lista.html',
                           pessoas=paginacao.items, paginacao=paginacao, q=q,
                           hoje=_date.today(),
                           filtro_tipo=filtro_tipo,
                           filtro_especial=filtro_especial,
                           filtro_tag=filtro_tag,
                           tags_disponiveis=tags_disponiveis,
                           ids_doc_incompleta=ids_doc_incompleta)


# ── Verificação de duplicidade (JSON) ─────────────────────────────────────────
@pessoas_bp.route('/verificar-duplicidade')
@login_required
def verificar_duplicidade():
    from flask import jsonify

    cpf        = request.args.get('cpf',  '').strip()
    cnpj       = request.args.get('cnpj', '').strip()
    excluir_id = request.args.get('excluir_id', type=int)

    duplicatas = []

    if cpf:
        q = Pessoa.visiveis().filter_by(cpf=cpf)
        if excluir_id:
            q = q.filter(Pessoa.id != excluir_id)
        for p in q.all():
            duplicatas.append({'id': p.id, 'nome': p.nome_razao_social,
                               'cpf': p.cpf or '', 'cnpj': p.cnpj or ''})

    if cnpj:
        ids_ja = {d['id'] for d in duplicatas}
        q = Pessoa.visiveis().filter_by(cnpj=cnpj)
        if excluir_id:
            q = q.filter(Pessoa.id != excluir_id)
        for p in q.all():
            if p.id not in ids_ja:
                duplicatas.append({'id': p.id, 'nome': p.nome_razao_social,
                                   'cpf': p.cpf or '', 'cnpj': p.cnpj or ''})

    return jsonify({'duplicatas': duplicatas})


# ── Exportar Excel ─────────────────────────────────────────────────────────────
@pessoas_bp.route('/exportar/xlsx')
@login_required
def exportar_xlsx():
    q               = request.args.get('q',       '').strip()
    filtro_tipo     = request.args.get('tipo',     '')
    filtro_especial = request.args.get('especial', '')
    query = Pessoa.visiveis()

    if q:
        like = f'%{q}%'
        query = query.filter(
            db.or_(
                Pessoa.nome_razao_social.ilike(like),
                Pessoa.cpf.ilike(like),
                Pessoa.cnpj.ilike(like),
            )
        )

    if filtro_tipo == 'PF':
        query = query.filter(
            Pessoa.cpf != None, Pessoa.cpf != '',
            db.or_(Pessoa.cnpj == None, Pessoa.cnpj == ''),
        )
    elif filtro_tipo == 'MEI':
        query = query.filter(Pessoa.mei == True)
    elif filtro_tipo == 'PJ':
        query = query.filter(
            Pessoa.cnpj != None, Pessoa.cnpj != '',
            db.or_(Pessoa.mei == None, Pessoa.mei == False),
        )

    if filtro_especial == 'sem_telefone':
        query = query.filter(db.or_(Pessoa.telefone == None, Pessoa.telefone == ''))
    elif filtro_especial == 'sem_atendimento':
        recentes = db.session.query(Atendimento.id_pessoa).filter(
            Atendimento.data_atendimento >= date.today() - timedelta(days=90),
            Atendimento.status == ATENDIMENTO_CONCLUIDO
        ).subquery()
        query = query.filter(~Pessoa.id.in_(recentes))

    pessoas = query.order_by(Pessoa.nome_razao_social).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pessoas"
    ws.sheet_view.showGridLines = False

    # Cabeçalho institucional
    ws.merge_cells('A1:G1')
    c = ws['A1']
    c.value = f"{marca.ORG_COMPLETO}"
    c.font = Font(bold=True, size=13, color="0f4c5c")
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 22

    ws.merge_cells('A2:G2')
    c2 = ws['A2']
    subtitulo = "Cadastro de Pessoas"
    if q:
        subtitulo += f' — Busca: "{q}"'
    subtitulo += f" · {len(pessoas)} registro(s)"
    c2.value = subtitulo
    c2.font = Font(size=9, color="64748B", italic=True)
    c2.alignment = Alignment(horizontal="center")

    ws.merge_cells('A3:G3')
    c3 = ws['A3']
    c3.value = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    c3.font = Font(size=8, color="94A3B8", italic=True)
    c3.alignment = Alignment(horizontal="center")

    # Linha de cabeçalho da tabela
    colunas = ["#", "Nome / Razão Social", "CPF", "CNPJ",
               "Telefone", "E-mail", "Endereço"]
    fill_cab = PatternFill("solid", fgColor="0f4c5c")
    borda = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for col_idx, nome in enumerate(colunas, start=1):
        cell = ws.cell(row=5, column=col_idx, value=nome)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = fill_cab
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = borda
    ws.row_dimensions[5].height = 18

    # Dados
    fill_par   = PatternFill("solid", fgColor="F1F5F9")
    fill_impar = PatternFill("solid", fgColor="FFFFFF")
    borda_dado = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for i, p in enumerate(pessoas, start=6):
        fill = fill_par if i % 2 == 0 else fill_impar
        valores = [
            p.id,
            p.nome_razao_social or "",
            p.cpf      or "",
            p.cnpj     or "",
            p.telefone or "",
            p.email    or "",
            p.endereco or "",
        ]
        for col_idx, val in enumerate(valores, start=1):
            cell = ws.cell(row=i, column=col_idx, value=val)
            cell.fill = fill
            cell.border = borda_dado
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center")

    # Larguras das colunas
    larguras = [6, 36, 16, 20, 16, 30, 40]
    for col_idx, larg in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = larg

    ws.freeze_panes = "A6"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    data_fmt = datetime.now().strftime('%d-%m-%Y')
    if q:
        nome_arquivo = f"Pessoas_{q[:20].replace(' ', '_')}_{data_fmt}.xlsx"
    else:
        nome_arquivo = f"Cadastro_de_Pessoas_{data_fmt}.xlsx"

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nome_arquivo,
    )


# ── Nova ───────────────────────────────────────────────────────────────────────
@pessoas_bp.route('/nova', methods=['GET', 'POST'])
@login_required
def nova():
    if request.method == 'POST':
        erros = _valida_form(request.form)
        if erros:
            for e in erros:
                flash(e, 'danger')
            return render_template('pessoas/form.html',
                                   titulo='Nova Pessoa', form=request.form, acao='nova',
                                   bairros=BAIRROS_MUNICIPIO)

        p = Pessoa()
        _salva_pessoa(p, request.form)
        from flask_login import current_user as _cu
        p.id_usuario_criador = _cu.id

        novo_aceite = request.form.get('aceite_whatsapp') in ('1', 'on', 'true')
        if novo_aceite:
            p.aceite_whatsapp = True
            p.data_aceite_whatsapp = datetime.now()
        else:
            p.aceite_whatsapp = False

        db.session.add(p)
        db.session.commit()
        flash(f'Pessoa "{p.nome_razao_social}" cadastrada com sucesso.', 'success')
        return redirect(url_for('pessoas.detalhe', id=p.id, iniciar_atendimento=1))

    tags_disponiveis = TagInteresse.query.filter_by(ativo=True).order_by(TagInteresse.descricao).all()
    return render_template('pessoas/form.html',
                           titulo='Nova Pessoa', form={}, acao='nova',
                           bairros=BAIRROS_MUNICIPIO,
                           tags_disponiveis=tags_disponiveis)


# ── Detalhe ────────────────────────────────────────────────────────────────────
@pessoas_bp.route('/<int:id>')
@login_required
def detalhe(id):
    pessoa = db.get_or_404(Pessoa, id)

    # Cadastro removido não tem ficha — é tela para agir sobre a pessoa
    if pessoa.removido_em is not None:
        flash('Este cadastro foi removido e não está mais disponível.', 'warning')
        return redirect(url_for('pessoas.index'))

    atendimentos = (
        Atendimento.query
        .filter_by(id_pessoa=id)
        .filter(Atendimento.status == ATENDIMENTO_CONCLUIDO)
        .order_by(Atendimento.data_atendimento.desc(),
                  Atendimento.hora_atendimento.desc())
        .all()
    )
    solicitacoes = (
        Solicitacao.query
        .filter_by(id_pessoa=id)
        .order_by(Solicitacao.data_solicitacao.desc())
        .all()
    )

    # ── Timeline unificada ────────────────────────────────────────────────
    from datetime import datetime as dt
    from flask import url_for as _url_for
    timeline = []
    for a in atendimentos:
        tipos_str = ', '.join(at.servico.descricao for at in a.servicos) or '—'
        timeline.append({
            'tipo':       'atendimento',
            'data':       dt.combine(a.data_atendimento, a.hora_atendimento),
            'data_fmt':   a.data_atendimento.strftime('%d/%m/%Y'),
            'hora_fmt':   a.hora_atendimento.strftime('%H:%M'),
            'responsavel': a.usuario.nome,
            'titulo':     'Atendimento',
            'subtitulo':  tipos_str,
            'detalhe':    a.observacao or '',
            'status':     None,
            'status_cls': None,
            'link':       _url_for('atendimentos.detalhe', id=a.id),
            'id':         a.id,
        })
    for s in solicitacoes:
        st = s.status.descricao.lower()
        if 'pendente' in st:
            cls = 'status-pendente'
        elif 'an' in st:
            cls = 'status-analise'
        elif 'aprovado' in st:
            cls = 'status-aprovado'
        elif 'recusado' in st:
            cls = 'status-recusado'
        else:
            cls = 'status-default'
        timeline.append({
            'tipo':       'solicitacao',
            'data':       dt.combine(s.data_solicitacao, dt.min.time()),
            'data_fmt':   s.data_solicitacao.strftime('%d/%m/%Y'),
            'hora_fmt':   None,
            'responsavel': s.usuario.nome,
            'titulo':     s.tipo_solicitacao.descricao,
            'subtitulo':  None,
            'detalhe':    s.descricao or '',
            'status':     s.status.descricao,
            'status_cls': cls,
            'link':       _url_for('solicitacoes.detalhe', id=s.id),
            'id':         s.id,
        })
    timeline.sort(key=lambda x: x['data'], reverse=True)

    # ── Resumo rápido ─────────────────────────────────────────────────────
    from datetime import date as _date
    hoje = _date.today()

    ultimo_atend = atendimentos[0] if atendimentos else None
    ultima_solic = solicitacoes[0] if solicitacoes else None

    solic_abertas = [
        s for s in solicitacoes
        if s.status.descricao.lower() in ('pendente', 'em análise', 'em analise')
    ]

    if ultimo_atend:
        dias_sem_atend = (hoje - ultimo_atend.data_atendimento).days
    else:
        dias_sem_atend = None

    resumo = {
        'total_atend':    len(atendimentos),
        'total_solic':    len(solicitacoes),
        'solic_abertas':  len(solic_abertas),
        'ultimo_atend':   ultimo_atend,
        'ultima_solic':   ultima_solic,
        'dias_sem_atend': dias_sem_atend,
    }

    from ...models import SolicitacaoExclusao
    exclusao_pendente = SolicitacaoExclusao.query.filter_by(
        entidade='cadastro',
        entidade_id=id,
        status='pendente'
    ).first()

    tags_disponiveis     = TagInteresse.query.filter_by(ativo=True).order_by(TagInteresse.descricao).all()
    tags_selecionadas_ids = {t.id for t in pessoa.tags}

    reqs = RequerimentoGerado.query.filter_by(id_pessoa=id).order_by(RequerimentoGerado.created_at).all()
    reqs_por_sol = {}
    for r in reqs:
        if r.id_solicitacao:
            reqs_por_sol.setdefault(r.id_solicitacao, []).append(r)

    return render_template('pessoas/detalhe.html',
                           pessoa=pessoa,
                           atendimentos=atendimentos,
                           solicitacoes=solicitacoes,
                           timeline=timeline,
                           resumo=resumo,
                           exclusao_pendente=exclusao_pendente,
                           tags_disponiveis=tags_disponiveis,
                           tags_selecionadas_ids=tags_selecionadas_ids,
                           reqs_por_sol=reqs_por_sol)


# ── Editar ─────────────────────────────────────────────────────────────────────
@pessoas_bp.route('/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def editar(id):
    pessoa = db.get_or_404(Pessoa, id)

    if pessoa.removido_em is not None:
        flash('Este cadastro foi removido e não pode ser editado.', 'warning')
        return redirect(url_for('pessoas.index'))

    if request.method == 'POST':
        erros = _valida_form(request.form)
        if erros:
            for e in erros:
                flash(e, 'danger')
            return render_template('pessoas/form.html',
                                   titulo='Editar Pessoa', form=request.form,
                                   acao='editar', pessoa=pessoa,
                                   bairros=BAIRROS_MUNICIPIO)

        aceite_anterior = pessoa.aceite_whatsapp
        _salva_pessoa(pessoa, request.form)

        novo_aceite = request.form.get('aceite_whatsapp') in ('1', 'on', 'true')
        if novo_aceite and not aceite_anterior:
            pessoa.aceite_whatsapp = True
            pessoa.data_aceite_whatsapp = datetime.now()
            pessoa.data_revogacao_whatsapp = None
        elif not novo_aceite and aceite_anterior:
            pessoa.aceite_whatsapp = False
            pessoa.data_revogacao_whatsapp = datetime.now()
        else:
            pessoa.aceite_whatsapp = novo_aceite

        db.session.commit()
        flash(f'Pessoa "{pessoa.nome_razao_social}" atualizada com sucesso.', 'success')
        return redirect(url_for('pessoas.detalhe', id=pessoa.id))

    # Pré-popula o form com os dados atuais
    form = {
        'nome_razao_social':       pessoa.nome_razao_social or '',
        'email':                   pessoa.email             or '',
        'cpf':                     pessoa.cpf               or '',
        'cnpj':                    pessoa.cnpj              or '',
        'data_nascimento':         pessoa.data_nascimento.isoformat() if pessoa.data_nascimento else '',
        'telefone':                pessoa.telefone          or '',
        'endereco':                pessoa.endereco          or '',
        'cep':                     pessoa.cep               or '',
        'logradouro':              pessoa.logradouro        or '',
        'numero':                  pessoa.numero            or '',
        'complemento':             pessoa.complemento       or '',
        'bairro':                  pessoa.bairro            or '',
        'sexo':                    pessoa.sexo              or '',
        'nome_fantasia':           pessoa.nome_fantasia          or '',
        'situacao_cadastral':      pessoa.situacao_cadastral     or '',
        'data_situacao_cadastral': pessoa.data_situacao_cadastral.isoformat() if pessoa.data_situacao_cadastral else '',
        'data_inicio_atividade':   pessoa.data_inicio_atividade.isoformat() if pessoa.data_inicio_atividade else '',
        'capital_social':          pessoa.capital_social         or '',
        'porte':                   pessoa.porte                  or '',
        'natureza_juridica':       pessoa.natureza_juridica      or '',
        'mei':                     'true' if pessoa.mei else ('false' if pessoa.mei is False else ''),
        'simples_nacional':        'true' if pessoa.simples_nacional else ('false' if pessoa.simples_nacional is False else ''),
        'cnae_principal':          pessoa.cnae_principal         or '',
        'cnae_secundario':         pessoa.cnae_secundario        or '',
        'inscricao_estadual':      pessoa.inscricao_estadual     or '',
        'rg':                      pessoa.rg                     or '',
        'nivel_escolar':           pessoa.nivel_escolar          or '',
        'num_membros_familia':     pessoa.num_membros_familia    if pessoa.num_membros_familia is not None else '',
        'renda_familiar':          pessoa.renda_familiar         or '',
        'instagram':               pessoa.instagram              or '',
        'endereco_comercial':      pessoa.endereco_comercial     or '',
        'principais_produtos':     pessoa.principais_produtos    or '',
        'faturamento_mensal':      pessoa.faturamento_mensal     or '',
        'num_funcionarios':        pessoa.num_funcionarios       if pessoa.num_funcionarios is not None else '',
        'observacoes_cadastro':    pessoa.observacoes_cadastro   or '',
        'aceite_whatsapp':         '1' if pessoa.aceite_whatsapp else '',
    }
    tags_disponiveis = TagInteresse.query.filter_by(ativo=True).order_by(TagInteresse.descricao).all()
    return render_template('pessoas/form.html',
                           titulo='Editar Pessoa', form=form,
                           acao='editar', pessoa=pessoa,
                           bairros=BAIRROS_MUNICIPIO,
                           tags_disponiveis=tags_disponiveis)


# ── Remover ────────────────────────────────────────────────────────────────────
_PERFIS_REMOVER = {'administrador'}


def _doc_str(pessoa):
    """Documento formatado para a mensagem de auditoria."""
    if pessoa.cpf and pessoa.cnpj:
        return f'CPF: {pessoa.cpf} / CNPJ: {pessoa.cnpj}'
    if pessoa.cnpj:
        return f'CNPJ: {pessoa.cnpj}'
    return f'CPF: {pessoa.cpf or "—"}'


def remover_cadastro(pessoa):
    """Exclusão lógica de cadastro — ponto único dos dois caminhos.

    Chamado por `pessoas.remover` (exclusão direta do admin) e por
    `pendencias.aprovar` (aprovação de solicitação de exclusão). Antes do
    ciclo 11 os dois divergiam: o primeiro bloqueava quando havia histórico,
    o segundo levantava IntegrityError ao tentar apagar.

    Marca a pessoa como removida e registra em LogAcao. **Nada é apagado** —
    atendimentos, solicitações, requerimentos e documentos permanecem
    intactos, e continuam contando nos relatórios.

    Não faz commit: quem chama decide o momento.
    Retorna False se o cadastro já estava removido.
    """
    if pessoa.removido_em is not None:
        return False

    pessoa.removido_em  = datetime.now(timezone.utc)
    pessoa.removido_por = current_user.id

    qtd_atend = Atendimento.query.filter_by(id_pessoa=pessoa.id).count()
    qtd_solic = Solicitacao.query.filter_by(id_pessoa=pessoa.id).count()

    _registrar_log(
        acao='exclusao',
        entidade='pessoa',
        entidade_id=pessoa.id,
        descricao=(
            f'Cadastro removido: {pessoa.nome_razao_social} ({_doc_str(pessoa)}) '
            f'— exclusão lógica; {qtd_atend} atendimento(s) e '
            f'{qtd_solic} solicitação(ões) preservados'
        ),
    )
    return True


@pessoas_bp.route('/<int:id>/remover', methods=['POST'])
@login_required
def remover(id):
    if current_user.perfil.nome not in _PERFIS_REMOVER:
        flash('Você não tem permissão para excluir cadastros.', 'danger')
        return redirect(url_for('pessoas.index'))

    pessoa = db.get_or_404(Pessoa, id)
    nome   = pessoa.nome_razao_social

    if not remover_cadastro(pessoa):
        flash('Este cadastro já foi removido.', 'warning')
        return redirect(url_for('pessoas.index'))

    db.session.commit()
    flash(f'Pessoa "{nome}" removida com sucesso. '
          f'O histórico de atendimentos foi preservado.', 'success')
    return redirect(url_for('pessoas.index'))


@pessoas_bp.route('/<int:id>/ficha-pdf')
@login_required
def ficha_pdf(id):
    """Gera PDF da ficha cadastral da pessoa."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table,
        TableStyle, HRFlowable, Image
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER

    pessoa       = db.get_or_404(Pessoa, id)
    atendimentos = (
        Atendimento.query
        .filter_by(id_pessoa=id)
        .filter(Atendimento.status == ATENDIMENTO_CONCLUIDO)
        .order_by(Atendimento.data_atendimento.desc())
        .all()
    )
    solicitacoes = (
        Solicitacao.query
        .filter_by(id_pessoa=id)
        .order_by(Solicitacao.data_solicitacao.desc())
        .all()
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm,
    )

    COR_PRIMARIA = colors.HexColor('#0f4c5c')
    COR_CINZA    = colors.HexColor('#f4f7f7')
    COR_TEXTO    = colors.HexColor('#0f2a30')
    BRANCO       = colors.white

    s_titulo = ParagraphStyle('titulo', fontSize=14, textColor=COR_PRIMARIA,
                               fontName='Helvetica-Bold', alignment=TA_CENTER, spaceAfter=4)
    s_sub    = ParagraphStyle('sub', fontSize=9, textColor=colors.HexColor('#6b7280'),
                               alignment=TA_CENTER, spaceAfter=12)
    s_secao  = ParagraphStyle('secao', fontSize=11, textColor=COR_PRIMARIA,
                               fontName='Helvetica-Bold', spaceBefore=14, spaceAfter=6)
    s_normal = ParagraphStyle('normal', fontSize=9, textColor=COR_TEXTO)

    story = []

    logo_path = os.path.join(current_app.root_path, 'static', 'img', 'logo.png')
    if os.path.exists(logo_path):
        from reportlab.lib.utils import ImageReader
        img_reader = ImageReader(logo_path)
        iw, ih = img_reader.getSize()
        ratio = ih / iw
        logo_w = 10*cm
        logo = Image(logo_path, width=logo_w, height=logo_w * ratio)
        logo.hAlign = 'CENTER'
        story.append(logo)
        story.append(Spacer(1, 6))

    story.append(Paragraph('Ficha Cadastral do Empreendedor', s_titulo))
    story.append(Paragraph(
        f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M")}',
        s_sub
    ))
    story.append(HRFlowable(width="100%", thickness=2, color=COR_PRIMARIA, spaceAfter=10))

    story.append(Paragraph('Dados Cadastrais', s_secao))

    tipo = 'Pessoa Física' if pessoa.cpf and not pessoa.cnpj else \
           'Pessoa Jurídica' if pessoa.cnpj else 'Não identificado'

    dados = [
        ['Nome / Razão Social', pessoa.nome_razao_social or '—'],
        ['Tipo',                tipo],
        ['CPF',                 pessoa.cpf  or '—'],
        ['CNPJ',                pessoa.cnpj or '—'],
        *([['CPF do Titular', pessoa.cpf]] if pessoa.cnpj and pessoa.cpf else []),
        ['Data de Nascimento',  pessoa.data_nascimento.strftime('%d/%m/%Y') if pessoa.data_nascimento else '—'],
        ['Telefone',            pessoa.telefone or '—'],
        ['E-mail',              pessoa.email    or '—'],
        ['Endereço', (
            ', '.join(filter(None, [
                pessoa.logradouro,
                f"Nº {pessoa.numero}" if pessoa.numero else None,
                pessoa.complemento,
                pessoa.bairro,
                f'{marca.CIDADE_UF}',
                f"CEP {pessoa.cep}" if pessoa.cep else None,
            ])) or pessoa.endereco or '—'
        )],
    ]

    # Dados da Receita Federal (apenas se preenchidos)
    dados_rf = []
    if pessoa.nome_fantasia:
        dados_rf.append(['Nome Fantasia', pessoa.nome_fantasia])
    if pessoa.situacao_cadastral:
        dados_rf.append(['Situação Cadastral', pessoa.situacao_cadastral])
    if pessoa.data_situacao_cadastral:
        dados_rf.append(['Data da Situação', pessoa.data_situacao_cadastral.strftime('%d/%m/%Y')])
    if pessoa.data_inicio_atividade:
        dados_rf.append(['Início da Atividade', pessoa.data_inicio_atividade.strftime('%d/%m/%Y')])
    if pessoa.porte:
        dados_rf.append(['Porte', pessoa.porte])
    if pessoa.natureza_juridica:
        dados_rf.append(['Natureza Jurídica', pessoa.natureza_juridica])
    if pessoa.capital_social:
        dados_rf.append(['Capital Social', pessoa.capital_social])
    if pessoa.inscricao_estadual:
        dados_rf.append(['Inscrição Estadual', pessoa.inscricao_estadual])
    if pessoa.mei is not None:
        dados_rf.append(['MEI', 'Sim' if pessoa.mei else 'Não'])
    if pessoa.simples_nacional is not None:
        dados_rf.append(['Simples Nacional', 'Sim' if pessoa.simples_nacional else 'Não'])
    if pessoa.cnae_principal:
        dados_rf.append(['CNAE Principal', pessoa.cnae_principal])
    if pessoa.cnae_secundario:
        dados_rf.append(['CNAEs Secundários', pessoa.cnae_secundario])

    t_dados = Table(dados, colWidths=[4.5*cm, 12.5*cm])
    t_dados.setStyle(TableStyle([
        ('FONTNAME',      (0,0), (0,-1), 'Helvetica-Bold'),
        ('FONTSIZE',      (0,0), (-1,-1), 9),
        ('TEXTCOLOR',     (0,0), (0,-1), COR_PRIMARIA),
        ('ROWBACKGROUNDS',(0,0), (-1,-1), [BRANCO, COR_CINZA]),
        ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
        ('TOPPADDING',    (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING',   (0,0), (-1,-1), 8),
    ]))
    story.append(t_dados)

    if dados_rf:
        story.append(Paragraph('Dados da Receita Federal', s_secao))
        t_rf = Table(dados_rf, colWidths=[4.5*cm, 12.5*cm])
        t_rf.setStyle(TableStyle([
            ('FONTNAME',      (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTSIZE',      (0,0), (-1,-1), 9),
            ('TEXTCOLOR',     (0,0), (0,-1), COR_PRIMARIA),
            ('ROWBACKGROUNDS',(0,0), (-1,-1), [BRANCO, COR_CINZA]),
            ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
            ('TOPPADDING',    (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING',   (0,0), (-1,-1), 8),
        ]))
        story.append(t_rf)

    dados_pessoais_comp = []
    if pessoa.sexo:
        sexo_label = 'Masculino' if pessoa.sexo == 'M' else 'Feminino' if pessoa.sexo == 'F' else pessoa.sexo
        dados_pessoais_comp.insert(0, ['Sexo', sexo_label])
    if pessoa.rg:
        dados_pessoais_comp.append(['RG', pessoa.rg])
    if pessoa.nivel_escolar:
        dados_pessoais_comp.append(['Nível Escolar', pessoa.nivel_escolar])
    if pessoa.num_membros_familia:
        dados_pessoais_comp.append(['Nº Membros da Família', str(pessoa.num_membros_familia)])
    if pessoa.renda_familiar:
        dados_pessoais_comp.append(['Renda Familiar', pessoa.renda_familiar])

    if dados_pessoais_comp:
        story.append(Paragraph('Dados Pessoais Complementares', s_secao))
        t_pc = Table(dados_pessoais_comp, colWidths=[4.5*cm, 12.5*cm])
        t_pc.setStyle(TableStyle([
            ('FONTNAME',      (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTSIZE',      (0,0), (-1,-1), 9),
            ('TEXTCOLOR',     (0,0), (0,-1), COR_PRIMARIA),
            ('ROWBACKGROUNDS',(0,0), (-1,-1), [BRANCO, COR_CINZA]),
            ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
            ('TOPPADDING',    (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING',   (0,0), (-1,-1), 8),
        ]))
        story.append(t_pc)

    dados_comerciais_comp = []
    if pessoa.instagram:
        dados_comerciais_comp.append(['Instagram', pessoa.instagram])
    if pessoa.faturamento_mensal:
        dados_comerciais_comp.append(['Faturamento Mensal', pessoa.faturamento_mensal])
    if pessoa.num_funcionarios is not None:
        dados_comerciais_comp.append(['Pessoas no Negócio', str(pessoa.num_funcionarios)])
    if pessoa.endereco_comercial:
        dados_comerciais_comp.append(['Endereço Comercial', pessoa.endereco_comercial])
    if pessoa.principais_produtos:
        dados_comerciais_comp.append(['Principais Produtos/Serviços', pessoa.principais_produtos])

    if dados_comerciais_comp:
        story.append(Paragraph('Dados Comerciais Complementares', s_secao))
        t_cc = Table(dados_comerciais_comp, colWidths=[4.5*cm, 12.5*cm])
        t_cc.setStyle(TableStyle([
            ('FONTNAME',      (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTSIZE',      (0,0), (-1,-1), 9),
            ('TEXTCOLOR',     (0,0), (0,-1), COR_PRIMARIA),
            ('ROWBACKGROUNDS',(0,0), (-1,-1), [BRANCO, COR_CINZA]),
            ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
            ('TOPPADDING',    (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING',   (0,0), (-1,-1), 8),
        ]))
        story.append(t_cc)

    if pessoa.observacoes_cadastro:
        story.append(Paragraph('Observações', s_secao))
        story.append(Paragraph(pessoa.observacoes_cadastro, s_normal))

    story.append(Paragraph(f'Atendimentos ({len(atendimentos)})', s_secao))

    if atendimentos:
        cab_at = [['Data', 'Hora', 'Atendente', 'Tipos', 'Observação']]
        rows_at = []
        for a in atendimentos:
            tipos_str = ', '.join(
                at.servico.descricao for at in a.servicos
            ) or '—'
            rows_at.append([
                a.data_atendimento.strftime('%d/%m/%Y'),
                a.hora_atendimento.strftime('%H:%M'),
                a.usuario.nome,
                tipos_str,
                (a.observacao or '—')[:60],
            ])
        t_at = Table(cab_at + rows_at, colWidths=[2.5*cm, 1.8*cm, 3*cm, 5*cm, 4.7*cm])
        t_at.setStyle(TableStyle([
            ('BACKGROUND',    (0,0), (-1,0), COR_PRIMARIA),
            ('TEXTCOLOR',     (0,0), (-1,0), BRANCO),
            ('FONTNAME',      (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',      (0,0), (-1,-1), 8),
            ('ROWBACKGROUNDS',(0,1), (-1,-1), [BRANCO, COR_CINZA]),
            ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
            ('TOPPADDING',    (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING',   (0,0), (-1,-1), 6),
        ]))
        story.append(t_at)
    else:
        story.append(Paragraph('Nenhum atendimento registrado.', s_normal))

    story.append(Paragraph(f'Solicitações ({len(solicitacoes)})', s_secao))

    if solicitacoes:
        cab_sol = [['Data', 'Tipo', 'Status', 'Responsável']]
        rows_sol = []
        for s in solicitacoes:
            rows_sol.append([
                s.data_solicitacao.strftime('%d/%m/%Y'),
                s.tipo_solicitacao.descricao[:35],
                s.status.descricao,
                s.usuario.nome,
            ])
        t_sol = Table(cab_sol + rows_sol, colWidths=[2.5*cm, 7*cm, 3*cm, 4.5*cm])
        t_sol.setStyle(TableStyle([
            ('BACKGROUND',    (0,0), (-1,0), COR_PRIMARIA),
            ('TEXTCOLOR',     (0,0), (-1,0), BRANCO),
            ('FONTNAME',      (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',      (0,0), (-1,-1), 8),
            ('ROWBACKGROUNDS',(0,1), (-1,-1), [BRANCO, COR_CINZA]),
            ('GRID',          (0,0), (-1,-1), 0.5, colors.HexColor('#dbe7e9')),
            ('TOPPADDING',    (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING',   (0,0), (-1,-1), 6),
        ]))
        story.append(t_sol)
    else:
        story.append(Paragraph('Nenhuma solicitação registrada.', s_normal))

    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1, color=COR_PRIMARIA))
    story.append(Paragraph(
        f'{marca.ORG_COMPLETO_UP} — ID #{pessoa.id}',
        ParagraphStyle('rodape', fontSize=7,
                       textColor=colors.HexColor('#9ca3af'),
                       alignment=TA_CENTER, spaceBefore=4)
    ))

    doc.build(story)
    buf.seek(0)

    nome_arquivo = f"ficha_{pessoa.nome_razao_social[:30].replace(' ', '_')}_{pessoa.id}.pdf"
    return send_file(
        buf,
        mimetype='application/pdf',
        as_attachment=False,
        download_name=nome_arquivo,
    )


# ── Sincronizar com Receita Federal ───────────────────────────────────────────
@pessoas_bp.route('/<int:id>/sincronizar-receita', methods=['POST'])
@login_required
def sincronizar_receita(id):
    from flask import jsonify
    import urllib.request
    import json as _json

    pessoa = db.get_or_404(Pessoa, id)

    if not pessoa.cnpj:
        return jsonify({'ok': False, 'erro': 'Cadastro não possui CNPJ.'}), 400

    cnpj_limpo = ''.join(filter(str.isdigit, pessoa.cnpj))
    if len(cnpj_limpo) != 14:
        return jsonify({'ok': False, 'erro': 'CNPJ inválido.'}), 400

    try:
        url = f'https://publica.cnpj.ws/cnpj/{cnpj_limpo}'
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = _json.loads(resp.read().decode())
    except Exception as e:
        return jsonify({'ok': False, 'erro': f'Erro ao consultar Receita Federal: {str(e)}'}), 502

    alteracoes = []

    def _atualizar(campo, novo_valor, label):
        if novo_valor is None:
            return
        novo_valor = str(novo_valor).strip() if not isinstance(novo_valor, bool) else novo_valor
        if not novo_valor:
            return
        atual = getattr(pessoa, campo)
        if str(atual or '').strip() != str(novo_valor).strip():
            alteracoes.append({
                'campo': label,
                'antes': str(atual or '—'),
                'depois': str(novo_valor)
            })
            setattr(pessoa, campo, novo_valor)

    est = data.get('estabelecimento', {})

    _atualizar('nome_fantasia',      est.get('nome_fantasia'),                    'Nome Fantasia')
    _atualizar('situacao_cadastral', est.get('situacao_cadastral'),               'Situação Cadastral')
    _atualizar('porte',              data.get('porte', {}).get('descricao'),       'Porte')
    _atualizar('natureza_juridica',  data.get('natureza_juridica', {}).get('descricao'), 'Natureza Jurídica')
    _atualizar('cnae_principal',     est.get('atividade_principal', {}).get('descricao'), 'CNAE Principal')

    for campo, chave, label in [
        ('data_situacao_cadastral', 'data_situacao_cadastral', 'Data da Situação'),
        ('data_inicio_atividade',   'data_inicio_atividade',   'Início da Atividade'),
    ]:
        val = est.get(chave)
        if val:
            try:
                from datetime import date as _date
                novo = _date.fromisoformat(val)
                atual = getattr(pessoa, campo)
                if atual != novo:
                    alteracoes.append({'campo': label,
                                       'antes': atual.strftime('%d/%m/%Y') if atual else '—',
                                       'depois': novo.strftime('%d/%m/%Y')})
                    setattr(pessoa, campo, novo)
            except Exception:
                pass

    if data.get('capital_social'):
        try:
            val_fmt = 'R$ ' + f"{float(data['capital_social']):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            _atualizar('capital_social', val_fmt, 'Capital Social')
        except Exception:
            pass

    mei_api = data.get('simples', {}).get('mei')
    if mei_api in ('Sim', 'Não'):
        novo_mei = mei_api == 'Sim'
        if pessoa.mei != novo_mei:
            alteracoes.append({'campo': 'MEI',
                               'antes': 'Sim' if pessoa.mei else 'Não' if pessoa.mei is False else '—',
                               'depois': 'Sim' if novo_mei else 'Não'})
            pessoa.mei = novo_mei

    sn_api = data.get('simples', {}).get('simples')
    if sn_api in ('Sim', 'Não'):
        novo_sn = sn_api == 'Sim'
        if pessoa.simples_nacional != novo_sn:
            alteracoes.append({'campo': 'Simples Nacional',
                               'antes': 'Sim' if pessoa.simples_nacional else 'Não' if pessoa.simples_nacional is False else '—',
                               'depois': 'Sim' if novo_sn else 'Não'})
            pessoa.simples_nacional = novo_sn

    cnaes_sec = est.get('atividades_secundarias', [])
    if cnaes_sec:
        novo_cnae_sec = '; '.join(a.get('descricao', '') for a in cnaes_sec if a.get('descricao'))
        _atualizar('cnae_secundario', novo_cnae_sec, 'CNAEs Secundários')

    inscricoes = est.get('inscricoes_estaduais', [])
    if inscricoes:
        ativa = next((i for i in inscricoes if i.get('ativo')), None)
        if ativa:
            _atualizar('inscricao_estadual', ativa.get('inscricao_estadual'), 'Inscrição Estadual')

    if not pessoa.telefone and est.get('ddd1') and est.get('telefone1'):
        num = est['telefone1'].replace('-', '').replace(' ', '')
        if len(num) == 8:
            num = '9' + num
        novo_tel = f"({est['ddd1']}) {num}"
        alteracoes.append({'campo': 'Telefone', 'antes': '—', 'depois': novo_tel})
        pessoa.telefone = novo_tel

    if not pessoa.email and est.get('email'):
        alteracoes.append({'campo': 'E-mail', 'antes': '—', 'depois': est['email']})
        pessoa.email = est['email']

    partes_end = [
        (est.get('tipo_logradouro', '') + ' ' + est.get('logradouro', '')).strip() or None,
        f"Nº {est['numero']}" if est.get('numero') else None,
        est.get('complemento') or None,
        est.get('bairro') or None,
        est.get('cidade', {}).get('nome') or None,
        est.get('estado', {}).get('sigla') or None,
        f"CEP {est['cep'][:5]}-{est['cep'][5:]}" if est.get('cep') and len(est['cep']) == 8 else None,
    ]
    novo_end = ', '.join(p for p in partes_end if p)
    if novo_end:
        _atualizar('endereco_comercial', novo_end, 'Endereço Comercial')

    if alteracoes:
        db.session.commit()
        _registrar_log(
            acao='sincronizacao',
            entidade='pessoa',
            entidade_id=id,
            descricao=f'Sincronização com Receita Federal: {len(alteracoes)} campo(s) atualizado(s) — CNPJ {pessoa.cnpj}'
        )

    return jsonify({
        'ok': True,
        'alteracoes': alteracoes,
        'total': len(alteracoes),
        'situacao': est.get('situacao_cadastral', ''),
    })


# ── Pendências rápidas (AJAX — usado em iniciar.html) ─────────────────────────
@pessoas_bp.route('/<int:id>/pendencias-rapido')
@login_required
def pendencias_rapido(id):
    from flask import jsonify
    from ...models import Status
    p = db.get_or_404(Pessoa, id)
    tipos_docs = {d.tipo for d in p.documentos_visiveis if d.ativo}
    tipos_obrigatorios = {'rg_cnh', 'comprovante_endereco'}
    doc_incompleta = not tipos_obrigatorios.issubset(tipos_docs)

    ids_abertos = [s.id for s in Status.query.filter(
        db.or_(
            Status.descricao.ilike('%pendente%'),
            Status.descricao.ilike('%análise%'),
        )
    ).all()]
    solic_aberta = Solicitacao.query.filter(
        Solicitacao.id_pessoa == id,
        Solicitacao.id_status.in_(ids_abertos)
    ).count() if ids_abertos else 0

    return jsonify({'doc_incompleta': doc_incompleta, 'solic_aberta': solic_aberta})


# ── Tags de Interesse ─────────────────────────────────────────────────────────
@pessoas_bp.route('/<int:id>/tags', methods=['POST'])
@login_required
def salvar_tags(id):
    pessoa = db.get_or_404(Pessoa, id)
    ids_selecionados = request.form.getlist('tags', type=int)
    tags_novas = TagInteresse.query.filter(TagInteresse.id.in_(ids_selecionados), TagInteresse.ativo == True).all()
    pessoa.tags = tags_novas
    db.session.commit()
    flash('Interesses atualizados com sucesso.', 'success')
    return redirect(url_for('pessoas.detalhe', id=id))


# ── Conversão PF → MEI ────────────────────────────────────────────────────────
@pessoas_bp.route('/<int:id>/converter-mei', methods=['POST'])
@login_required
def converter_mei(id):
    pessoa = db.get_or_404(Pessoa, id)

    cnpj = request.form.get('cnpj', '').strip()
    if not cnpj:
        return jsonify({'ok': False, 'erro': 'CNPJ é obrigatório.'}), 422

    if not _CNPJ_RE.match(cnpj):
        return jsonify({'ok': False, 'erro': 'CNPJ fora do formato 00.000.000/0000-00.'}), 422
    if not _cnpj_valido(cnpj):
        return jsonify({'ok': False, 'erro': 'CNPJ inválido — verifique os dígitos.'}), 422

    pessoa.mei  = True
    pessoa.cnpj = cnpj
    _registrar_log(
        acao='conversao_mei',
        entidade='pessoa',
        entidade_id=id,
        descricao=f'Conversão PF → MEI: CNPJ {cnpj} — {pessoa.nome_razao_social}',
    )
    db.session.commit()
    return jsonify({'ok': True})


# ── LGPD — Consentimento WhatsApp ─────────────────────────────────────────────
@pessoas_bp.route('/<int:id>/whatsapp-consentimento', methods=['POST'])
@login_required
def whatsapp_consentimento(id):
    pessoa = db.get_or_404(Pessoa, id)
    acao = request.form.get('acao', '')

    if acao == 'conceder' and not pessoa.aceite_whatsapp:
        pessoa.aceite_whatsapp = True
        pessoa.data_aceite_whatsapp = datetime.now()
        pessoa.data_revogacao_whatsapp = None
        _registrar_log(
            acao='lgpd_aceite',
            entidade='pessoa',
            entidade_id=id,
            descricao=f'Consentimento WhatsApp concedido — {pessoa.nome_razao_social}',
        )
        db.session.commit()
        flash('Consentimento WhatsApp registrado.', 'success')
    elif acao == 'revogar' and pessoa.aceite_whatsapp:
        pessoa.aceite_whatsapp = False
        pessoa.data_revogacao_whatsapp = datetime.now()
        _registrar_log(
            acao='lgpd_revogacao',
            entidade='pessoa',
            entidade_id=id,
            descricao=f'Consentimento WhatsApp revogado — {pessoa.nome_razao_social}',
        )
        db.session.commit()
        flash('Consentimento WhatsApp revogado.', 'success')

    return redirect(url_for('pessoas.detalhe', id=id))
