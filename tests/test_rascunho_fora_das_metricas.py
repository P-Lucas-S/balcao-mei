"""
test_rascunho_fora_das_metricas.py — Rascunho não é número oficial.

Desde que o ícone de "Solicitação" passou a salvar o atendimento como
rascunho, existem `Atendimento` com `status='rascunho'` de verdade. Um
atendimento aberto e abandonado **não pode contar como atendimento feito**.

A prova aqui não é ler o código: é medir. Cada tela de gestão é renderizada,
um rascunho é inserido, e a tela é renderizada de novo. Se um único número
mudou, algum SELECT esqueceu o filtro — e o teste diz qual tela.

É a mesma classe de cuidado do soft-delete de pessoa: esconder em tudo que é
número oficial, sem exceção.
"""
import re
from datetime import date

from app import db as _db
from app.models import (Atendimento, ATENDIMENTO_CONCLUIDO,
                        ATENDIMENTO_RASCUNHO)
from tests.conftest import AtendimentoFactory, PessoaFactory


# Toda tela de gestão que exibe contagem de atendimento.
TELAS_DE_GESTAO = [
    "/dashboard",
    "/analise/",
    "/analise/jornada",
    "/analise/demanda",
    "/analise/perfil",
    "/analise/potencial",
    "/relatorios/",
    "/relatorios/atendimentos",
    "/relatorios/desempenho",
    "/relatorios/executivo",
]


def _numeros(html):
    """Só os números da página — o que uma métrica vazada mudaria.

    Descarta o CSRF (muda a cada request) e os blocos <script>, que carregam
    tokens e dados de gráfico com ruído próprio.
    """
    html = re.sub(r"<script.*?</script>", "", html, flags=re.S)
    html = re.sub(r'name="csrf_token"[^>]*', "", html)
    # O indicador "Em aberto" DEVE mudar com um rascunho — é o que ele mede.
    # Ele fica fora da comparação justamente por não ser número oficial.
    html = re.sub(r'<section id="indicador-em-aberto".*?</section>', "", html, flags=re.S)
    return re.findall(r"\d+", html)


def _linhas_xlsx(rv):
    """Linhas por aba de uma planilha exportada — {aba: nº de linhas}."""
    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(rv.get_data()), read_only=True)
    return {ws.title: ws.max_row for ws in wb.worksheets}


def _atendimento(usuario, pessoa, status, dia=None):
    a = AtendimentoFactory.create(
        id_pessoa=pessoa.id, id_usuario=usuario.id, status=status,
        data_atendimento=dia or date.today(),
    )
    _db.session.commit()
    return a


class TestRascunhoNaoContaEmNenhumaTela:
    """A varredura: nenhuma tela de gestão pode mexer por causa de um rascunho."""

    def test_nenhum_numero_muda_ao_inserir_rascunho(self, client, db, auth):
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        # Uma base concluída, para as telas terem o que mostrar
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)

        antes = {}
        for url in TELAS_DE_GESTAO:
            rv = client.get(url)
            assert rv.status_code == 200, f"{url} devolveu {rv.status_code}"
            antes[url] = _numeros(rv.get_data(as_text=True))

        _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        vazaram = []
        for url in TELAS_DE_GESTAO:
            depois = _numeros(client.get(url).get_data(as_text=True))
            if depois != antes[url]:
                vazaram.append(url)

        assert not vazaram, (
            "o rascunho vazou para os números destas telas: "
            + ", ".join(vazaram))

    def test_atendimento_concluido_conta_normalmente(self, client, db, auth):
        """O contrapeso: se nada mudasse nunca, o teste acima seria vazio."""
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)

        antes = _numeros(client.get("/dashboard").get_data(as_text=True))
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)
        depois = _numeros(client.get("/dashboard").get_data(as_text=True))

        assert depois != antes, "atendimento concluído tem de mexer no número"

    def test_exportacoes_tambem_ignoram_rascunho(self, client, db, auth):
        """Relatório exportado é registro oficial — mesma regra da tela.

        A comparação é por LINHAS da planilha, não por tamanho do arquivo:
        .xlsx é um zip e carrega timestamp, então dois downloads iguais podem
        ter bytes diferentes. Contar bytes daria um teste instável — e foi
        exatamente o que ele fez antes desta correção.
        """
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)

        urls = ["/relatorios/atendimentos/exportar/xlsx",
                "/relatorios/desempenho/exportar/xlsx",
                "/relatorios/executivo/exportar/xlsx"]
        antes = {}
        for url in urls:
            rv = client.get(url)
            assert rv.status_code == 200, f"{url} devolveu {rv.status_code}"
            antes[url] = _linhas_xlsx(rv)

        _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        for url in urls:
            assert _linhas_xlsx(client.get(url)) == antes[url], (
                f"o rascunho entrou na exportação {url}")


class TestContagensDiretas:
    """Alguns números merecem asserção explícita, não só 'não mudou'."""

    def _kpis(self, client):
        html = client.get("/dashboard").get_data(as_text=True)
        return _numeros(html)

    def test_dashboard_conta_so_concluido(self, client, db, auth):
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        for _ in range(3):
            _atendimento(u, p, ATENDIMENTO_CONCLUIDO)
        for _ in range(5):
            _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        html = client.get("/dashboard").get_data(as_text=True)
        assert Atendimento.query.count() == 8
        # 3 concluídos hoje — o número do card, não os 8 do banco
        assert re.search(r">\s*3\s*<", html), "esperava o card mostrando 3"

    def test_ficha_da_pessoa_nao_lista_rascunho(self, client, db, auth):
        """A ficha não desenha o selo, então rascunho ali passaria por feito."""
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)

        def _timeline():
            html = client.get(f"/pessoas/{p.id}").get_data(as_text=True)
            return re.findall(r'count-badge">(\d+)', html)

        antes = _timeline()
        _atendimento(u, p, ATENDIMENTO_RASCUNHO)
        assert _timeline() == antes, "o rascunho entrou na timeline da ficha"

    def test_pessoa_com_so_rascunho_conta_como_sem_atendimento(self, client, db, auth):
        """O filtro 'sem atendimento há 90 dias' não se satisfaz com rascunho."""
        u = auth.login_como("atendente")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        html = client.get("/pessoas/?filtro=sem_atendimento").get_data(as_text=True)
        assert p.nome_razao_social in html


class TestIndicadorEmAberto:
    """O indicador operacional — separado do número oficial, nunca somado."""

    def test_analise_mostra_quantos_estao_em_aberto(self, client, db, auth):
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)
        for _ in range(2):
            _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        html = client.get("/analise/").get_data(as_text=True)
        bloco = re.search(r'id="indicador-em-aberto".*?</section>', html, re.S)
        assert bloco, "a análise deve expor o indicador de atendimentos em aberto"
        bloco = bloco.group(0)

        # O total, no título — e não qualquer "2" perdido numa data.
        assert re.search(r"2\s+atendimentos?\s+pendentes?", bloco), (
            "o indicador deve dizer quantos estão em aberto")
        assert re.search(r'painel-em-aberto-qtd">2<', bloco), (
            "a quebra por atendente deve mostrar os 2 do mesmo usuário")

    def test_indicador_some_quando_nao_ha_rascunho(self, client, db, auth):
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_CONCLUIDO)

        html = client.get("/analise/").get_data(as_text=True)
        assert 'id="indicador-em-aberto"' not in html

    def test_indicador_diz_que_nao_e_metrica_oficial(self, client, db, auth):
        """A tela precisa dizer que são coisas distintas, não só mostrar."""
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        html = client.get("/analise/").get_data(as_text=True)
        bloco = re.search(r'id="indicador-em-aberto".*?</section>', html, re.S).group(0)
        assert "não entram" in bloco or "não contam" in bloco

    def test_indicador_quebra_por_atendente(self, client, db, auth):
        u = auth.login_como("administrador")
        p = PessoaFactory.create()
        _atendimento(u, p, ATENDIMENTO_RASCUNHO)

        html = client.get("/analise/").get_data(as_text=True)
        bloco = re.search(r'id="indicador-em-aberto".*?</section>', html, re.S).group(0)
        assert u.nome in bloco
